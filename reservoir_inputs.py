"""Read the small, numeric Excel reservoir specification used by 03 and 04.

The workbook owns reservoir inputs. This module only validates them and maps
their explicit column units to the ordinary dictionaries used by ESBMTK.
"""

from __future__ import annotations

import math
from numbers import Real
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_RESERVOIR_WORKBOOK = (
    Path(__file__).resolve().parent / "data/Boudreau_2010/model_definition.xlsx"
)
# Preserve the benchmark's object/equation ordering even after sorting Excel.
BOX_ORDER = ("H_b", "L_b", "D_b")
OCEAN_COLUMNS = (
    "Box ID", "Description", "Area (m2)", "Volume (m3)",
    "Temperature (degC)", "Salinity", "Pressure (bar)",
    "Initial DIC (umol/kg)", "Initial TA (umol/kg)",
)
ATMOSPHERE_COLUMNS = (
    "Box ID", "Total air (mol)", "Initial CO2 (ppm)", "Role",
)


def read_reservoir_tables(path=DEFAULT_RESERVOIR_WORKBOOK):
    """Return named-table records, retaining their visible Excel column labels.

    Require literal inputs rather than potentially stale cached formula results.
    Tables can move on the sheet and their rows/columns can be reordered.
    """
    return _read_tables(path, {
        'OceanReservoirs': ('Reservoirs', OCEAN_COLUMNS),
        'Atmosphere': ('Reservoirs', ATMOSPHERE_COLUMNS),
    })


def _read_tables(path, specifications):
    with Path(path).open("rb") as stream:
        workbook = load_workbook(stream, data_only=False)
        try:
            result = {}
            for name, (sheet_name, columns) in specifications.items():
                if sheet_name not in workbook:
                    raise ValueError(f"Missing Excel sheet {sheet_name!r}")
                sheet = workbook[sheet_name]
                if name not in sheet.tables:
                    raise ValueError(f"Missing Excel table {name!r} on {sheet_name}")
                cells = sheet[sheet.tables[name].ref]
                headers = [cell.value for cell in cells[0]]
                if len(headers) != len(columns) or set(headers) != set(columns):
                    raise ValueError(f"{name}: keep the column names and units: {columns}")
                records = []
                for row in cells[1:]:
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError(f"{name}: use literal input values, not formulas")
                    records.append(dict(zip(headers, (cell.value for cell in row))))
                result[name] = records
            return result
        finally:
            workbook.close()


def _number(record, column, *, minimum=None, maximum=None, positive=False):
    value = record[column]
    label = f"{record['Box ID']} / {column}"
    if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value):
        raise ValueError(f"{label}: enter a finite number, not a blank or text")
    if positive and value <= 0:
        raise ValueError(f"{label}: must be positive")
    if minimum is not None and value < minimum:
        raise ValueError(f"{label}: must be at least {minimum}")
    if maximum is not None and value > maximum:
        raise ValueError(f"{label}: must be at most {maximum}")
    return float(value)


def load_reservoir_inputs(path=DEFAULT_RESERVOIR_WORKBOOK):
    """Validate and translate ocean and atmosphere inputs; create no objects."""
    tables = read_reservoir_tables(path)
    records = tables["OceanReservoirs"]
    ids = [row["Box ID"] for row in records]
    if len(ids) != len(BOX_ORDER) or set(ids) != set(BOX_ORDER):
        raise ValueError(f"OceanReservoirs needs each Box ID exactly once: {BOX_ORDER}")
    by_id = {row["Box ID"]: row for row in records}
    boxes = {}
    for name in BOX_ORDER:
        row = by_id[name]
        if not isinstance(row["Description"], str) or not row["Description"].strip():
            raise ValueError(f"{name}: add a descriptive name")
        boxes[name] = {
            "dic": f"{_number(row, 'Initial DIC (umol/kg)', positive=True):.17g} umol/kg",
            "ta": f"{_number(row, 'Initial TA (umol/kg)', minimum=0):.17g} umol/kg",
            "area": f"{_number(row, 'Area (m2)', positive=True):.17g} m**2",
            "volume": f"{_number(row, 'Volume (m3)', positive=True):.17g} m**3",
            "temperature": _number(row, "Temperature (degC)", minimum=-2, maximum=40),
            "pressure": _number(row, "Pressure (bar)", minimum=0, maximum=1100),
            "salinity": _number(row, "Salinity", minimum=0, maximum=50),
        }
    air = tables["Atmosphere"]
    if len(air) != 1 or air[0]["Box ID"] != "CO2_At":
        raise ValueError("Atmosphere needs exactly one row with Box ID CO2_At")
    return {
        "boxes": boxes,
        "pco2": f"{_number(air[0], 'Initial CO2 (ppm)', positive=True, maximum=1e6):.17g} ppm",
        "atmosphere_moles": f"{_number(air[0], 'Total air (mol)', positive=True):.17g} mol",
    }


def reservoir_inventory_rows(model):
    """Display ESBMTK-derived mass and its current initial dissolved inventories.

    Call before a run. If read_state was used, index zero is the restart state.
    TA amounts are reported as mol equivalents; ESBMTK stores them as mol.
    """
    rows = []
    for name in BOX_ORDER:
        box = getattr(model, name)
        volume = box.DIC.volume.to("m**3").magnitude
        density = box.swc.density
        mass = volume * density
        rows.append({
            "Box": name,
            "Density (kg/m3)": density,
            "Water mass (kg)": mass,
            "Initial carbon (mol)": mass * box.DIC.c[0],
            "Initial TA (mol eq)": mass * box.TA.c[0],
        })
    return rows

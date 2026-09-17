"""Check Excel input handling and the resulting ESBMTK inventories."""

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
from esbmtk import Q_

from model import initialize_model, run_model
from presets import load_boudreau_parameters, make_process_variant
from reservoir_inputs import (
    BOX_ORDER, DEFAULT_RESERVOIR_WORKBOOK, load_reservoir_inputs,
    read_reservoir_tables, reservoir_inventory_rows,
)


class ReservoirWorkbookTest(unittest.TestCase):
    def test_archived_geometry_and_initial_inventory_are_preserved(self):
        p = load_boudreau_parameters()
        # Independent migration/regression checks, not editable model defaults.
        expected = {
            'H_b': (0.5e14, 1.76e16, 2, 35, 17.6, 2153, 2345),
            'L_b': (2.85e14, 2.85e16, 21.5, 35, 5, 1952, 2288),
            'D_b': (3.36e14, 1.29e18, 2, 35, 240, 2291, 2399),
        }
        for name, values in expected.items():
            b = p['boxes'][name]
            actual = (Q_(b['area']).magnitude, Q_(b['volume']).magnitude,
                      b['temperature'], b['salinity'], b['pressure'],
                      Q_(b['dic']).magnitude, Q_(b['ta']).magnitude)
            self.assertEqual(actual, values)
        self.assertEqual(Q_(p['atmosphere_moles']).magnitude, 1.7786e20)
        self.assertEqual(Q_(p['pco2']).magnitude, 280)

    def test_editing_workbook_changes_actual_model_inputs_without_rounding(self):
        # Patch numeric cells in a disposable XLSX to exercise the real reader.
        edits = {'E8': 3e16, 'I8': 1952.123456789, 'C15': 1.8e20, 'D15': 300}
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'edited.xlsx'
            with ZipFile(DEFAULT_RESERVOIR_WORKBOOK) as src, ZipFile(path, 'w') as dst:
                for item in src.infolist():
                    data = src.read(item.filename)
                    if item.filename == 'xl/worksheets/sheet1.xml':
                        sheet = ET.fromstring(data)
                        for address, value in edits.items():
                            cell = sheet.find(f".//s:c[@r='{address}']/s:v", ns)
                            self.assertIsNotNone(cell)
                            cell.text = str(value)
                        data = ET.tostring(sheet, encoding='utf-8')
                    dst.writestr(item, data)
            p = load_boudreau_parameters(path)
            model = initialize_model(p, stop='1 yr', max_timestep='1 yr')
            np.testing.assert_allclose(model.L_b.DIC.volume.to('m**3').magnitude, 3e16, rtol=1e-14)
            self.assertAlmostEqual(model.L_b.DIC.c[0] * 1e6, edits['I8'], places=9)
            self.assertEqual(model.CO2_At.v[0], 1.8e20)
            self.assertAlmostEqual(model.CO2_At.c[0] * 1e6, 300)

    def test_sorting_rows_does_not_change_construction_order(self):
        tables = read_reservoir_tables()
        tables['OceanReservoirs'].reverse()
        with patch('reservoir_inputs.read_reservoir_tables', return_value=tables):
            self.assertEqual(tuple(load_reservoir_inputs()['boxes']), BOX_ORDER)

    def test_invalid_values_and_duplicate_or_missing_boxes_fail(self):
        for column, value in (
            ('Volume (m3)', None), ('Volume (m3)', 0), ('Area (m2)', -1),
            ('Initial DIC (umol/kg)', '1952'), ('Salinity', float('nan')),
            ('Temperature (degC)', 100), ('Pressure (bar)', -5),
        ):
            with self.subTest(column=column, value=value):
                tables = read_reservoir_tables()
                tables['OceanReservoirs'][0][column] = value
                with patch('reservoir_inputs.read_reservoir_tables', return_value=tables):
                    with self.assertRaises(ValueError):
                        load_reservoir_inputs()
        for ids in (['H_b', 'H_b', 'D_b'], ['H_b', 'L_b']):
            tables = read_reservoir_tables()
            tables['OceanReservoirs'] = tables['OceanReservoirs'][:len(ids)]
            for row, name in zip(tables['OceanReservoirs'], ids):
                row['Box ID'] = name
            with patch('reservoir_inputs.read_reservoir_tables', return_value=tables):
                with self.assertRaisesRegex(ValueError, 'exactly once'):
                    load_reservoir_inputs()

    def test_reader_rejects_formula_and_changed_unit_header(self):
        # Mutations remain in memory; no test authoring changes the source XLSX.
        from openpyxl import load_workbook
        for address, value, message in (
            ('I7', '=2000+153', 'literal'),
            ('H6', 'Pressure (dbar)', 'column names and units'),
        ):
            with self.subTest(address=address):
                wb = load_workbook(DEFAULT_RESERVOIR_WORKBOOK)
                wb['Reservoirs'][address] = value
                with patch('reservoir_inputs.load_workbook', return_value=wb):
                    with self.assertRaisesRegex(ValueError, message):
                        read_reservoir_tables()
                wb.close()

    def test_atmosphere_requires_one_valid_finite_inventory(self):
        for key, value in (('Total air (mol)', 0), ('Initial CO2 (ppm)', None),
                           ('Initial CO2 (ppm)', 1e7), ('Box ID', 'ocean')):
            tables = read_reservoir_tables()
            tables['Atmosphere'][0][key] = value
            with patch('reservoir_inputs.read_reservoir_tables', return_value=tables):
                with self.assertRaises(ValueError):
                    load_reservoir_inputs()

    def test_models_use_esbmtk_density_and_report_implemented_inventories(self):
        model = initialize_model(stop='1 yr', max_timestep='1 yr')
        run_model(model)
        for row in reservoir_inventory_rows(model):
            box = getattr(model, row['Box'])
            # Legacy pre-run .m arrays use nominal litres; the ODE coefficients
            # are the independent check of the mass used by the actual solver.
            for species, label in ((box.DIC, 'Initial carbon (mol)'), (box.TA, 'Initial TA (mol eq)')):
                nonzero = model.CM[species.idx]
                nonzero = nonzero[nonzero != 0]
                np.testing.assert_allclose(abs(nonzero), 1 / row['Water mass (kg)'], rtol=1e-12)
                np.testing.assert_allclose(row[label], species.c[0] / abs(nonzero[0]), rtol=1e-12)

    def test_fresh_loads_and_matched_cases_are_independent(self):
        p = load_boudreau_parameters()
        original = deepcopy(p)
        case = make_process_variant(base=p)
        case['boxes']['L_b']['dic'] = '1900 umol/kg'
        self.assertEqual(p, original)
        p['boxes']['H_b']['temperature'] = 5
        self.assertEqual(load_boudreau_parameters(), original)


if __name__ == '__main__':
    unittest.main()

"""Check the paper-to-workbook cross-reference and student worksheet masking."""
import unittest
from pathlib import Path

from teaching_specification import flux_specification, flux_table_markdown, workbook_connections


class TeachingSpecificationTest(unittest.TestCase):
    def test_connection_keys_disambiguate_shared_circulation_id(self):
        rows = workbook_connections()
        ids = [r['ID'] for r in rows]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(sum(i.endswith('@thc') for i in ids), 3)

    def test_student_records_exclude_selected_answers(self):
        rows = {r['Process ID']: r for r in flux_specification(student=True)}
        for row in rows.values():
            self.assertEqual(row['J_TA (eq/yr)'], 'Your expression')
            if row['Process ID'] != 'D':
                self.assertEqual(row['J_DIC (mol C/yr)'], 'Your expression')
        self.assertEqual(rows['POC']['Process / arrow'], 'Organic export: your arrow')
        self.assertIn('supplied', rows['D']['J_DIC (mol C/yr)'])
        self.assertNotIn('B_{net}(t)=E(t)-D(t)', flux_table_markdown(student=True))

    def test_exported_student_file_contains_no_hidden_answer_sheet_or_formulas(self):
        from openpyxl import load_workbook
        root = Path(__file__).resolve().parents[1]
        path = root / 'outputs/03_04_flux_specification/student.xlsx'
        with path.open('rb') as stream:
            wb = load_workbook(stream, data_only=False)
            self.assertEqual(wb.sheetnames, ['FluxSpecification', 'WorkbookLinks'])
            self.assertTrue(all(s.sheet_state == 'visible' for s in wb))
            self.assertTrue(all(c.data_type != 'f' for s in wb for row in s for c in row))
            sheet = wb['FluxSpecification']
            # POC arrow; paired carbon/TA expressions, with dissolution supplied.
            for address in ('C9', 'D7', 'E7', 'D10', 'E10', 'E12', 'D13', 'E13'):
                self.assertIsNone(sheet[address].value)
            self.assertIn('D(t)', sheet['D12'].value)
            self.assertEqual(sheet['D6'].value, 'J_DIC (mol C/yr)')
            self.assertEqual(sheet['E6'].value, 'J_TA (eq/yr)')
            wb.close()
